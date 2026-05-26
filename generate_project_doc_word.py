"""
Converts BlueLine_Project_Documentation.md to a professionally formatted Word document.
Run: python generate_project_doc_word.py
Output: BlueLine_Project_Documentation.docx
"""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INPUT_FILE  = Path(__file__).parent / "BlueLine_Project_Documentation.md"
OUTPUT_FILE = Path(__file__).parent / "BlueLine_Project_Documentation.docx"

# Colour palette
BLUE_DARK   = RGBColor(0x1A, 0x3A, 0x6B)
BLUE_MID    = RGBColor(0x1E, 0x5F, 0xAF)
BLUE_LIGHT  = RGBColor(0x2E, 0x86, 0xC1)
TEAL        = RGBColor(0x00, 0x7A, 0x87)
CODE_BG     = RGBColor(0xF6, 0xF8, 0xFA)
CODE_FG     = RGBColor(0x24, 0x29, 0x2E)


def set_paragraph_shading(paragraph, fill_hex: str):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    pPr.append(shd)


def set_cell_shading(cell, fill_hex: str):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tcPr.append(shd)


def set_cell_border(cell, color="C8D4E3"):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"),   "single")
        tag.set(qn("w:sz"),    "4")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)
        tcBorders.append(tag)
    tcPr.append(tcBorders)


def add_page_number(doc: Document):
    section = doc.sections[0]
    footer  = section.footer
    para    = footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    para.clear()

    run = para.add_run("Page ")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    for field_name, run_obj in [("PAGE", para.add_run()), ("NUMPAGES", para.add_run())]:
        if field_name == "NUMPAGES":
            r = para.add_run(" of ")
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        fldChar  = OxmlElement("w:fldChar")
        fldChar.set(qn("w:fldCharType"), "begin")
        instrText = OxmlElement("w:instrText")
        instrText.text = field_name
        fldChar2 = OxmlElement("w:fldChar")
        fldChar2.set(qn("w:fldCharType"), "end")
        run_obj.font.size = Pt(9)
        run_obj.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        run_obj._r.append(fldChar)
        run_obj._r.append(instrText)
        run_obj._r.append(fldChar2)


def add_title_page(doc: Document):
    for _ in range(5):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("PROJECT BLUELINE")
    r.font.size = Pt(34)
    r.font.bold = True
    r.font.color.rgb = BLUE_DARK

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = subtitle.add_run("Project Documentation")
    r2.font.size = Pt(20)
    r2.font.color.rgb = BLUE_MID

    tagline = doc.add_paragraph()
    tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = tagline.add_run("AI-Powered Automation — Quality Gate  |  Security Loop  |  Certificate Loop")
    r3.font.size = Pt(11)
    r3.font.italic = True
    r3.font.color.rgb = TEAL

    doc.add_paragraph()
    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r4 = line.add_run("─" * 42)
    r4.font.color.rgb = BLUE_LIGHT
    doc.add_paragraph()

    for label, value in [
        ("Version",     "1.0"),
        ("Date",        "2026-05-26"),
        ("Status",      "In Progress"),
        ("Prepared by", "Project BlueLine Team"),
    ]:
        meta = doc.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = meta.add_run(f"{label}: ")
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = BLUE_DARK
        r2 = meta.add_run(value)
        r2.font.size = Pt(11)
        r2.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    doc.add_page_break()


def apply_inline(para, text: str):
    parts = re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            r = para.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = para.add_run(part[1:-1])
            r.font.name  = "Consolas"
            r.font.size  = Pt(9)
            r.font.color.rgb = RGBColor(0xC7, 0x25, 0x4F)
        else:
            if part:
                para.add_run(part)


def parse_table(doc: Document, table_lines: list):
    rows = []
    for line in table_lines:
        if re.match(r'^\s*\|[-| :]+\|\s*$', line):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if cells:
            rows.append(cells)

    if not rows:
        return

    col_count = max(len(r) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=col_count)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT

    for i, row_data in enumerate(rows):
        for j, cell_text in enumerate(row_data):
            if j >= col_count:
                break
            cell = tbl.cell(i, j)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            para = cell.paragraphs[0]
            para.clear()
            cell_text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', cell_text)

            if i == 0:
                set_cell_shading(cell, "1A3A6B")
                r = para.add_run(cell_text)
                r.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                r.font.size = Pt(9)
            else:
                set_cell_shading(cell, "FFFFFF" if i % 2 == 0 else "EBF5FB")
                apply_inline(para, cell_text)
                for run in para.runs:
                    run.font.size = Pt(9)

            set_cell_border(cell)

    doc.add_paragraph()


def parse_code_block(doc: Document, lines: list, lang: str):
    if lang and lang.strip():
        lbl = doc.add_paragraph()
        lbl.paragraph_format.space_before = Pt(4)
        lbl.paragraph_format.space_after  = Pt(0)
        r = lbl.add_run(f"  {lang.strip()}")
        r.font.size = Pt(8)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        set_paragraph_shading(lbl, "DEE3EA")

    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(0)
        p.paragraph_format.left_indent  = Cm(0.5)
        set_paragraph_shading(p, "F6F8FA")
        r = p.add_run(line.rstrip())
        r.font.name  = "Consolas"
        r.font.size  = Pt(8)
        r.font.color.rgb = CODE_FG

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after  = Pt(6)
    set_paragraph_shading(spacer, "F6F8FA")


def add_horizontal_rule(doc: Document):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after  = Pt(6)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"),   "single")
    bottom.set(qn("w:sz"),    "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "2E86C1")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_agent_section_divider(doc: Document, agent_name: str, track: str, agent_id: str):
    doc.add_page_break()
    banner = doc.add_paragraph()
    banner.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_shading(banner, "1A3A6B")
    r1 = banner.add_run(f"  {agent_name}")
    r1.font.size = Pt(18)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    sub = doc.add_paragraph()
    set_paragraph_shading(sub, "2E5F9A")
    r2 = sub.add_run(f"  {track}   |   ID: {agent_id}")
    r2.font.size = Pt(10)
    r2.font.color.rgb = RGBColor(0xCC, 0xDD, 0xFF)
    doc.add_paragraph()


def convert_md_to_docx(md_path: Path, out_path: Path):
    doc = Document()

    section = doc.sections[0]
    section.page_width    = Inches(8.5)
    section.page_height   = Inches(11)
    section.left_margin   = Inches(1.0)
    section.right_margin  = Inches(1.0)
    section.top_margin    = Inches(1.0)
    section.bottom_margin = Inches(1.0)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    add_page_number(doc)
    add_title_page(doc)

    lines = md_path.read_text(encoding="utf-8").splitlines()

    # Agent metadata for banner dividers
    AGENT_META = {
        "Agent 1: CLARION":    ("Quality Gate Track", "BL-QG-001"),
        "Agent 2: LUMEN":      ("Quality Gate Track", "BL-QG-002"),
        "Agent 3: VECTOR":     ("Quality Gate Track", "BL-QG-003"),
        "Agent 4: ASCENT":     ("Quality Gate Track", "BL-QG-004"),
        "Agent 5: BULWARK":    ("Security Track",     "BL-SEC-001"),
        "Agent 6: WATCHTOWER": ("Security Track",     "BL-SEC-002"),
        "Agent 7: FORGE":      ("Security Track",     "BL-SEC-003"),
        "Agent 8: STEWARD":    ("Security Track",     "BL-SEC-004"),
        "Agent 9: TIMELINE":   ("Certificate Loop Track", "BL-CERT-001"),
        "Agent 10: REGENT":    ("Certificate Loop Track", "BL-CERT-002"),
        "Agent 11: COURIER":   ("Certificate Loop Track", "BL-CERT-003"),
        "Agent 12: HARBOUR":   ("Certificate Loop Track", "BL-CERT-004"),
    }

    i = 0
    table_lines = []
    code_lines  = []
    in_code     = False
    code_lang   = ""
    in_table    = False

    while i < len(lines):
        line = lines[i]

        # Code block
        if line.startswith("```"):
            if not in_code:
                in_code   = True
                code_lang = line[3:].strip()
                code_lines = []
            else:
                parse_code_block(doc, code_lines, code_lang)
                in_code   = False
                code_lang = ""
                code_lines = []
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        # Table detection
        is_table_line = line.strip().startswith("|") and "|" in line[1:]

        if is_table_line:
            in_table = True
            table_lines.append(line)
            i += 1
            continue
        elif in_table:
            parse_table(doc, table_lines)
            table_lines = []
            in_table    = False

        # Horizontal rule
        if re.match(r'^-{3,}$', line.strip()):
            add_horizontal_rule(doc)
            i += 1
            continue

        # H1 — document title (skip, already on title page)
        if re.match(r'^# [^#]', line):
            if "Project BlueLine" in line and "Documentation" in line:
                i += 1
                continue
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after  = Pt(6)
            r = p.add_run(line[2:].strip())
            r.font.size = Pt(20)
            r.font.bold = True
            r.font.color.rgb = BLUE_DARK
            i += 1
            continue

        # H2 — section headings
        if line.startswith("## "):
            heading_text = line[3:].strip()

            # Agent section — insert coloured banner
            matched_agent = None
            for key in AGENT_META:
                if key in heading_text:
                    matched_agent = key
                    break

            if matched_agent:
                track, agent_id = AGENT_META[matched_agent]
                agent_name = matched_agent.split(": ", 1)[1]
                add_agent_section_divider(doc, agent_name, track, agent_id)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after  = Pt(4)
                r = p.add_run(heading_text)
                r.font.size = Pt(15)
                r.font.bold = True
                r.font.color.rgb = BLUE_MID
                pPr  = p._p.get_or_add_pPr()
                pBdr = OxmlElement("w:pBdr")
                btm  = OxmlElement("w:bottom")
                btm.set(qn("w:val"),   "single")
                btm.set(qn("w:sz"),    "4")
                btm.set(qn("w:space"), "1")
                btm.set(qn("w:color"), "2E86C1")
                pBdr.append(btm)
                pPr.append(pBdr)
            i += 1
            continue

        # H3
        if line.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after  = Pt(3)
            r = p.add_run(line[4:].strip())
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = BLUE_LIGHT
            i += 1
            continue

        # H4
        if line.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after  = Pt(2)
            r = p.add_run(line[5:].strip())
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            i += 1
            continue

        # Bullet list
        if re.match(r'^(\s*)[-*] ', line):
            indent = len(line) - len(line.lstrip())
            text   = re.sub(r'^(\s*)[-*] ', '', line)
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent  = Cm(0.5 + indent * 0.2)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after  = Pt(1)
            p.clear()
            bullet_run = p.add_run("• ")
            bullet_run.font.color.rgb = BLUE_MID
            bullet_run.font.bold      = True
            apply_inline(p, text.strip())
            for r in p.runs[1:]:
                r.font.size = Pt(10)
            i += 1
            continue

        # Numbered list
        if re.match(r'^\d+\. ', line):
            text = re.sub(r'^\d+\. ', '', line)
            num  = re.match(r'^(\d+)\.', line).group(1)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent  = Cm(0.5)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after  = Pt(1)
            nr = p.add_run(f"{num}. ")
            nr.font.bold      = True
            nr.font.color.rgb = BLUE_MID
            apply_inline(p, text.strip())
            i += 1
            continue

        # Block quote
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent  = Cm(1.0)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after  = Pt(4)
            set_paragraph_shading(p, "EBF5FB")
            apply_inline(p, line[2:].strip())
            for r in p.runs:
                r.font.size      = Pt(9.5)
                r.font.italic    = True
                r.font.color.rgb = RGBColor(0x1A, 0x5C, 0x85)
            i += 1
            continue

        # Blank line
        if not line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after  = Pt(3)
            i += 1
            continue

        # Skip metadata lines already on title page
        if re.match(r'^\*\*(Version|Date|Status|Prepared by)\*\*', line):
            i += 1
            continue

        # "Agent-Level Documentation Checklist" divider heading
        if line.strip() == "# Agent-Level Documentation Checklist":
            doc.add_page_break()
            p = doc.add_paragraph()
            set_paragraph_shading(p, "1A3A6B")
            r = p.add_run("  Agent-Level Documentation Checklist")
            r.font.size = Pt(18)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            i += 1
            continue

        # Regular paragraph
        stripped = line.strip()
        if re.match(r'^[=*]{3,}$', stripped):
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(4)
        stripped = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', stripped)
        apply_inline(p, stripped)
        i += 1

    if table_lines:
        parse_table(doc, table_lines)

    doc.save(out_path)
    print(f"[OK] Saved: {out_path}")
    print(f"     Lines processed: {len(lines)}")


if __name__ == "__main__":
    print("Converting BlueLine_Project_Documentation.md to .docx ...")
    convert_md_to_docx(INPUT_FILE, OUTPUT_FILE)
